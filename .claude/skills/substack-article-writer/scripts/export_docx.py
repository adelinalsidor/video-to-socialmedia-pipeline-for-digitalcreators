#!/usr/bin/env python3
"""
Convertește article.md (Markdown cu frontmatter YAML) într-un .docx simplu,
lizibil de client fără unelte tehnice. Titlul devine H1 în document (Substack
îl afișează automat ca H1, deci nu apare ca `#` în body-ul .md), `##`/`###`
devin Heading 2/3, `**bold**` devine bold. Subtitle-ul din frontmatter apare
ca paragraf-introducere italic sub titlu.

Usage:
    python3 export_docx.py --in article.md --out article.docx
"""

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def parse_frontmatter(text: str):
    if not text.startswith("---"):
        return {}, text
    end = text.find("\n---", 3)
    if end == -1:
        return {}, text
    raw = text[3:end].strip()
    body = text[end + 4:].lstrip("\n")

    meta = {}
    for line in raw.splitlines():
        if ":" not in line:
            continue
        key, _, value = line.partition(":")
        meta[key.strip()] = value.strip().strip('"')
    return meta, body


def add_runs_with_bold(paragraph, text: str):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def convert(md_path: Path, docx_path: Path):
    text = md_path.read_text(encoding="utf-8")
    meta, body = parse_frontmatter(text)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)

    title = meta.get("title", md_path.stem)
    doc.add_heading(title, level=1)

    subtitle = meta.get("subtitle")
    if subtitle:
        p = doc.add_paragraph()
        run = p.add_run(subtitle)
        run.italic = True

    for raw_line in body.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        else:
            p = doc.add_paragraph()
            add_runs_with_bold(p, line)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"Gata: {docx_path}")
    print(f"\nDOCX_PATH={docx_path}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--in", dest="in_path", required=True)
    parser.add_argument("--out", dest="out_path", required=True)
    args = parser.parse_args()

    in_path = Path(args.in_path)
    if not in_path.exists():
        print(f"ERROR: nu există {in_path}", file=sys.stderr)
        sys.exit(1)

    convert(in_path, Path(args.out_path))


if __name__ == "__main__":
    main()
